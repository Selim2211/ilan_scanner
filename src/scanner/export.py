"""Excel / CSV cikti."""
from __future__ import annotations

import csv
import io
import json
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MODE_LABEL = {"remote": "UZAKTAN", "hybrid": "Hibrit", "onsite": "Yerinde", "unknown": "?"}

COLUMNS = [
    ("Skor", "score", 8),
    ("Calisma", "work_mode", 12),
    ("Sozlesme", "is_contract", 12),
    ("Baslik", "title", 45),
    ("Firma", "company", 26),
    ("Konum", "location", 20),
    ("Sure", "duration", 14),
    ("Baslangic", "starts_at", 12),
    ("Butce/Ucret", "budget_raw", 18),
    ("Ilan tarihi", "posted_at", 14),
    ("Kaynak", "source", 14),
    ("Eslesen", "keywords_hit", 30),
    ("Link", "url", 55),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
REMOTE_FILL = PatternFill("solid", fgColor="C6EFCE")
HYBRID_FILL = PatternFill("solid", fgColor="FFEB9C")


def _json_list(value) -> str:
    if not value:
        return ""
    try:
        return ", ".join(json.loads(value))
    except (json.JSONDecodeError, TypeError):
        return str(value)


def _value(row: sqlite3.Row, key: str) -> str:
    value = row[key]
    if key == "work_mode":
        return MODE_LABEL.get(value or "unknown", "?")
    if key == "is_contract":
        return "" if value is None else ("Proje/Contract" if value else "Kadrolu")
    if key in ("skills", "keywords_hit"):
        return _json_list(value)
    if key == "posted_at" and value:
        return str(value)[:10]
    return "" if value is None else str(value)


def to_xlsx(rows: Sequence[sqlite3.Row], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SAP Projeleri"

    for col, (header, _key, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(rows, start=2):
        for c, (_header, key, _w) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=_value(row, key))
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("title", "company"))
            if key == "url" and row["url"]:
                cell.hyperlink = row["url"]
                cell.font = Font(color="0563C1", underline="single")
        fill = REMOTE_FILL if row["work_mode"] == "remote" else (
            HYBRID_FILL if row["work_mode"] == "hybrid" else None)
        if fill:
            ws.cell(row=r, column=2).fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"
    wb.save(path)
    return path


def to_csv(rows: Sequence[sqlite3.Row], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow([header for header, _k, _w in COLUMNS])
        for row in rows:
            writer.writerow([_value(row, key) for _h, key, _w in COLUMNS])
    return path


def default_name(prefix: str, extension: str) -> str:
    return f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M')}.{extension}"


# --- panel "Dışa aktar" ekrani: secili ilanlar, okunur rapor -----------------

STATUS_LABEL = {
    "new": "İşaretsiz", "shortlist": "Takipte", "applied": "Başvuruldu",
    "rejected": "Olumsuz", "won_pending": "Olumlu — başlamadı",
    "won_active": "Olumlu — çalışılıyor", "won_done": "Olumlu — bitti",
}

#: (baslik, genislik, sarma) - deger _report_value ile uretilir
REPORT_COLUMNS = [
    ("Güncellik", 14, False),
    ("Skor", 7, False),
    ("Başlık", 46, True),
    ("İlan tarihi", 13, False),
    ("Çalışma şekli", 13, False),
    ("Firma", 24, True),
    ("Kaynak", 13, False),
    ("Bütçe / Ücret", 16, False),
    ("Durum", 16, False),
    ("Önemli hususlar", 60, True),
    ("Bağlantı", 50, False),
]


def _guncellik(row: sqlite3.Row) -> str:
    if not row["is_active"]:
        return "İlan kapandı"
    if (row["missing_streak"] or 0) >= 2:
        return "Kontrol ediliyor"
    return "Yayında"


def _onemli_hususlar(row: sqlite3.Row) -> str:
    """AI'nin cikardigi 'olmazsa olmaz' maddeler; yoksa eslesen kelimeler."""
    maddeler = _json_list_items(row["ai_must_haves"] if "ai_must_haves" in row.keys() else None)
    if maddeler:
        return "\n".join(f"• {m}" for m in maddeler)
    hits = _json_list_items(row["keywords_hit"])
    return ", ".join(hits)


def _json_list_items(value) -> list[str]:
    if not value:
        return []
    try:
        return [str(x).strip() for x in json.loads(value) if str(x).strip()]
    except (json.JSONDecodeError, TypeError):
        return []


def _report_value(row: sqlite3.Row, header: str) -> str:
    if header == "Güncellik":
        return _guncellik(row)
    if header == "Skor":
        return "" if row["score"] is None else str(row["score"])
    if header == "Başlık":
        return row["title"] or ""
    if header == "İlan tarihi":
        return str(row["posted_at"] or row["first_seen_at"] or "")[:10]
    if header == "Çalışma şekli":
        return MODE_LABEL.get(row["work_mode"] or "unknown", "?")
    if header == "Firma":
        return row["company"] or ""
    if header == "Kaynak":
        return row["source"] or ""
    if header == "Bütçe / Ücret":
        return row["budget_raw"] or ""
    if header == "Durum":
        mark = row["mark"] if "mark" in row.keys() else ""
        return STATUS_LABEL.get(mark, "")
    if header == "Önemli hususlar":
        return _onemli_hususlar(row)
    if header == "Bağlantı":
        return row["url"] or ""
    return ""


def report_xlsx_bytes(rows: Sequence[sqlite3.Row]) -> bytes:
    """Secili ilanlari okunur, bicimli bir Excel'e yazar; baytlari dondurur (panel indirir)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Seçili ilanlar"

    for col, (header, width, _wrap) in enumerate(REPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 20

    for r, row in enumerate(rows, start=2):
        for c, (header, _w, wrap) in enumerate(REPORT_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=_report_value(row, header))
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if header == "Bağlantı" and row["url"]:
                cell.hyperlink = row["url"]
                cell.font = Font(color="0563C1", underline="single")
        mod = row["work_mode"]
        if mod == "remote":
            ws.cell(row=r, column=5).fill = REMOTE_FILL
        elif mod == "hybrid":
            ws.cell(row=r, column=5).fill = HYBRID_FILL
        if not row["is_active"]:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="F2C6C6")

    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_COLUMNS))}{max(ws.max_row, 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
