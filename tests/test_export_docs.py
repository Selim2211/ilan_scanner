"""Word (.docx) ve Excel (.xlsx) çıktı üreticilerinin birim testleri.

Panel rotalarından bağımsız: `docx_export.build_docx` ve `export.report_xlsx_bytes`
doğrudan sözlük/satır listeleriyle çağrılır. Ağ yok, DB yok.
"""
import io
import zipfile
from xml.dom.minidom import parseString

from openpyxl import load_workbook

from scanner.docx_export import build_docx
from scanner.export import report_xlsx_bytes


def _row(**over):
    base = dict(
        fingerprint="fp1", title="SAP ABAP Danışmanı", score=72, is_active=1,
        work_mode="remote", is_contract=1, company="Acme GmbH", location="Berlin",
        country="DE", duration="6 ay", starts_at="10/2026", budget_raw="€600/gün",
        posted_at="2026-09-01T00:00:00+00:00", first_seen_at="2026-09-02T00:00:00+00:00",
        source="freelancermap", url="https://x/1", mark="applied", missing_streak=0,
        ai_summary="", ai_must_haves=None, summary_en=None,
        keywords_hit='["abap","sap","remote"]', description="Uzun açıklama metni.",
    )
    base.update(over)
    return base


# --- docx -----------------------------------------------------------------

def _docx_xml(rows):
    data = build_docx(rows)
    z = zipfile.ZipFile(io.BytesIO(data))
    assert set(z.namelist()) >= {"[Content_Types].xml", "_rels/.rels",
                                 "word/document.xml", "word/styles.xml"}
    doc = z.read("word/document.xml").decode("utf-8")
    parseString(doc)                       # iyi biçimli XML (yoksa fırlatır)
    return doc


def test_docx_kunye_alanlari_yazilir():
    doc = _docx_xml([_row()])
    for beklenen in ["SAP ABAP Danışmanı", "Künye", "72", "Acme GmbH", "freelancermap",
                     "€600/gün", "Uzaktan", "Başvuruldu", "https://x/1"]:
        assert beklenen in doc, beklenen


def test_docx_kosul_yoksa_eslesen_kelimeler_bolumu():
    doc = _docx_xml([_row(ai_must_haves=None)])
    assert "Olmazsa olmaz koşullar" not in doc
    assert "Eşleşen kelimeler" in doc and "abap, sap, remote" in doc


def test_docx_kosul_varsa_madde_listesi():
    doc = _docx_xml([_row(ai_summary="Kısa özet.",
                          ai_must_haves='["5+ yıl ABAP", "Akıcı İngilizce"]')])
    assert "Özet" in doc and "Kısa özet." in doc
    assert "Olmazsa olmaz koşullar" in doc
    assert "5+ yıl ABAP" in doc and "Akıcı İngilizce" in doc


def test_docx_ozel_karakter_xml_kacisi():
    doc = _docx_xml([_row(title='SAP & <ABAP> "Remote" Danışman',
                          company="Acme & Co <GmbH>")])
    assert "&amp;" in doc and "&lt;ABAP&gt;" in doc
    assert "<ABAP>" not in doc                       # ham açı ayracı kalmamalı


def test_docx_kapali_ilan_guncellik_metni():
    doc = _docx_xml([_row(is_active=0)])
    assert "İlan kapandı" in doc


def test_docx_iki_ilan_sayfa_sonu():
    doc = _docx_xml([_row(fingerprint="a", title="Birinci"),
                     _row(fingerprint="b", title="İkinci")])
    assert 'w:type="page"' in doc                    # aralarına sayfa sonu
    assert "Birinci" in doc and "İkinci" in doc


def test_docx_bos_baslik_dusmez():
    doc = _docx_xml([_row(title=None)])
    assert "başlıksız ilan" in doc


# --- xlsx ---------------------------------------------------------------

def _sheet(rows):
    wb = load_workbook(io.BytesIO(report_xlsx_bytes(rows)))
    return wb.active


def test_xlsx_basliklar_ve_sutun_sirasi():
    ws = _sheet([_row()])
    basliklar = [c.value for c in ws[1]]
    assert basliklar[:4] == ["Güncellik", "Skor", "Başlık", "İlan tarihi"]
    assert "Önemli hususlar" in basliklar and "Bağlantı" in basliklar


def test_xlsx_guncellik_degerleri():
    ws = _sheet([
        _row(fingerprint="a", is_active=1, missing_streak=0),
        _row(fingerprint="b", is_active=1, missing_streak=3),
        _row(fingerprint="c", is_active=0),
    ])
    guncellik = [ws.cell(row=r, column=1).value for r in (2, 3, 4)]
    assert guncellik == ["Yayında", "Kontrol ediliyor", "İlan kapandı"]


def test_xlsx_onemli_hususlar_kaynagi():
    # ai_must_haves varsa madde madde
    ws = _sheet([_row(ai_must_haves='["Senior ABAP", "S/4HANA"]')])
    hus_col = [c.value for c in ws[1]].index("Önemli hususlar") + 1
    hucre = ws.cell(row=2, column=hus_col).value
    assert "Senior ABAP" in hucre and "•" in hucre
    # yoksa eşleşen kelimeler
    ws2 = _sheet([_row(ai_must_haves=None)])
    hucre2 = ws2.cell(row=2, column=hus_col).value
    assert hucre2 == "abap, sap, remote"


def test_xlsx_baglanti_hyperlink():
    ws = _sheet([_row(url="https://ornek/ilan/9")])
    link_col = [c.value for c in ws[1]].index("Bağlantı") + 1
    assert ws.cell(row=2, column=link_col).hyperlink is not None


def test_xlsx_satir_sayisi_baslik_dahil():
    ws = _sheet([_row(fingerprint=f"f{i}") for i in range(4)])
    assert ws.max_row == 5                           # 4 ilan + başlık
