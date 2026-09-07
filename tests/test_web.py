"""Panel rotalari: ayarlar kaydi, filtre parametrelerinin sayfalamada korunmasi."""
import pytest
from fastapi.testclient import TestClient

from scanner import config as config_mod
from scanner.scheduler import BackgroundScanner


@pytest.fixture
def anon_client(tmp_path, monkeypatch):
    """Oturum acmamis istemci - giris duvarini sinayan testler icin."""
    monkeypatch.setattr(config_mod.paths, "data_dir", lambda: tmp_path)
    from scanner.web.app import create_app
    with TestClient(create_app(auto_scan=False)) as c:
        yield c


@pytest.fixture
def client(anon_client):
    """Panel gecici bir veri klasorunde calisir: gercek DB ve overlay'e dokunulmaz.

    Panel artik giris istiyor; ilk acilista kurulan yonetici hesabiyla oturum acilir,
    boylece rota testleri yetki duvarina takilmadan asil isi sinar.
    """
    anon_client.post("/giris", data={"username": "admin", "password": "123456", "next": "/"},
                     follow_redirects=False)
    yield anon_client


def test_ayarlar_sayfasi_kaynaklari_listeler(client):
    response = client.get("/ayarlar")
    assert response.status_code == 200
    assert "Ayarlar" in response.text
    assert "jooble" in response.text
    assert "Veriyi kontrol et" in response.text


def test_tarama_ayari_overlaye_yazilir(client, tmp_path):
    response = client.post("/ayarlar/tarama", data={
        "auto_enabled": "1", "interval_minutes": "25", "quick_interval_minutes": "7",
        "verify_limit": "50", "verify_every_n_scans": "1", "page_size": "35",
    }, follow_redirects=False)

    assert response.status_code == 303
    assert response.headers["location"].startswith("/ayarlar?")
    data = config_mod.load_overlay(config_mod.overlay_path())
    assert data["auto_scan"]["interval_minutes"] == 25
    assert data["panel"]["page_size"] == 35


def test_tarama_araligi_alt_sinirin_altina_inmez(client):
    client.post("/ayarlar/tarama", data={"interval_minutes": "1", "quick_interval_minutes": "1"},
                follow_redirects=False)
    data = config_mod.load_overlay(config_mod.overlay_path())
    assert data["auto_scan"]["interval_minutes"] == 5
    assert data["auto_scan"]["quick_interval_minutes"] == 2


TEMEL_ARAMA = {"queries": "SAP ABAP", "must_any": "abap\nsap"}


def test_ne_ariyoruz_bozuk_agirlikta_hata_dondurur(client):
    response = client.post("/ayarlar/ne-ariyoruz", data={**TEMEL_ARAMA, "boost": "abap: cok"},
                           follow_redirects=False)
    assert response.status_code == 303
    assert "hata=" in response.headers["location"]


def test_ne_ariyoruz_bos_zorunlu_kelimeyi_reddeder(client):
    """must_any bosalirsa hicbir ilan elenmez ve liste alakasiz ilanla dolar."""
    response = client.post("/ayarlar/ne-ariyoruz", data={"queries": "ERP", "must_any": ""},
                           follow_redirects=False)
    assert "hata=" in response.headers["location"]
    assert not config_mod.keywords_overlay_path().exists()


def test_ne_ariyoruz_konuyu_degistirir(client):
    """SAP -> ERP: sorgular ve zorunlu kelimeler tek formdan degisir."""
    response = client.post("/ayarlar/ne-ariyoruz", data={
        "queries": "ERP yazilim\nERP developer",
        "must_any": "erp",
        "boost": "erp: 8",
        "min_score": "5",
    }, follow_redirects=False)
    assert response.status_code == 303
    assert "ok=" in response.headers["location"]

    assert config_mod.load_config()["queries"] == ["ERP yazilim", "ERP developer"]
    keywords = config_mod.load_keywords()
    assert keywords["must_any"] == ["erp"]
    assert keywords["boost"] == {"erp": 8}


def test_ne_ariyoruz_ic_ice_sinyali_korur(client):
    """signals.country_boost formda gosterilemiyor; kaydederken silinmemeli."""
    onceki = config_mod.load_keywords()["signals"]["country_boost"]
    client.post("/ayarlar/ne-ariyoruz", data={
        **TEMEL_ARAMA, "signals": "remote_boost: 20\ncontract_boost: 8"},
        follow_redirects=False)

    signals = config_mod.load_keywords()["signals"]
    assert signals["remote_boost"] == 20
    assert signals["country_boost"] == onceki      # dokunulmadi
    assert "hybrid_boost" not in signals           # formdan silinen gitti


def test_ne_ariyoruz_ekrandan_okunani_geri_yazmak_bozmaz(client):
    """Sayfayi acip hicbir sey degistirmeden kaydetmek ayarlari bozmamali."""
    from scanner import settings as settings_mod

    once = config_mod.load_keywords()
    client.post("/ayarlar/ne-ariyoruz", data={
        "queries": "\n".join(config_mod.load_config()["queries"]),
        "must_any": settings_mod.format_lines(once["must_any"]),
        "hard_exclude": settings_mod.format_lines(once["hard_exclude"]),
        "boost": settings_mod.format_weight_lines(once["boost"]),
        "penalty": settings_mod.format_weight_lines(once["penalty"]),
        "signals": settings_mod.format_weight_lines(once["signals"]),
        "min_score": str(once["min_score"]),
        "title_multiplier": str(once["title_multiplier"]),
    }, follow_redirects=False)

    assert config_mod.load_keywords() == once


# --- kodsuz yeni kaynak -----------------------------------------------
YENI_KAYNAK = {
    "title": "Deneme Panosu", "kind": "json",
    "url": "https://ornek.test/api?q={sorgu}",
    "items_path": "data", "pages": "1",
    "field__title": "baslik", "field__url": "link", "enabled": "1",
}


def test_yeni_kaynak_eklenir_ve_taramaya_girer(client):
    response = client.post("/ayarlar/kaynak-ekle", data=YENI_KAYNAK, follow_redirects=False)
    assert response.status_code == 303
    assert "ok=" in response.headers["location"]

    config = config_mod.load_config()
    kaynak = config["sources"]["deneme-panosu"]
    assert kaynak["kind"] == "json"
    assert kaynak["fields"] == {"title": "baslik", "url": "link"}

    # kodsuz kaynak gercekten tarama listesine giriyor mu
    from scanner.pipeline import active_sources
    assert "deneme-panosu" in active_sources(config)


def test_yeni_kaynak_gecersiz_adresi_reddeder(client):
    response = client.post("/ayarlar/kaynak-ekle",
                           data={**YENI_KAYNAK, "url": "ornek.test/api"},
                           follow_redirects=False)
    assert "hata=" in response.headers["location"]
    assert "deneme-panosu" not in (config_mod.load_config().get("sources") or {})


def test_yeni_kaynak_json_alan_eslemesi_ister(client):
    response = client.post("/ayarlar/kaynak-ekle",
                           data={**YENI_KAYNAK, "field__title": "", "field__url": ""},
                           follow_redirects=False)
    assert "hata=" in response.headers["location"]


def test_yeni_kaynak_kodda_tanimli_adi_kullanamaz(client):
    response = client.post("/ayarlar/kaynak-ekle", data={**YENI_KAYNAK, "title": "jooble"},
                           follow_redirects=False)
    assert "hata=" in response.headers["location"]


def test_ozel_kaynak_silinir(client):
    client.post("/ayarlar/kaynak-ekle", data=YENI_KAYNAK, follow_redirects=False)
    response = client.post("/ayarlar/kaynak-sil", data={"name": "deneme-panosu"},
                           follow_redirects=False)
    assert "ok=" in response.headers["location"]
    assert "deneme-panosu" not in (config_mod.load_config().get("sources") or {})


def test_kodda_tanimli_kaynak_silinemez(client):
    response = client.post("/ayarlar/kaynak-sil", data={"name": "jooble"},
                           follow_redirects=False)
    assert "hata=" in response.headers["location"]
    assert "jooble" in config_mod.load_config()["sources"]


def test_anahtar_kaydi_bos_alani_yok_sayar(client):
    response = client.post("/ayarlar/anahtarlar", data={"env__JOOBLE_API_KEY": ""},
                           follow_redirects=False)
    assert response.status_code == 303
    assert not config_mod.env_path().exists()      # bos form dosya olusturmamali


def test_sayfalama_baglantisi_yeni_filtreleri_tasir(client, tmp_path):
    """base_query'ye eklenmezse tarih araligi 2. sayfada sessizce kaybolurdu."""
    from datetime import datetime, timedelta, timezone

    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    # sayfalama baglantisi ancak birden fazla sayfa varken cikar
    client.post("/ayarlar/tarama", data={"page_size": "5"}, follow_redirects=False)
    store = Storage(tmp_path / "data" / "projects.db")
    bugun = datetime.now(timezone.utc)
    ilanlar = []
    for i in range(12):
        p = Project(source="test", source_id=str(i), url=f"http://x/{i}",
                    title=f"SAP ABAP {i}", company=f"F{i}",
                    posted_at=bugun - timedelta(days=3))
        p.fingerprint = fingerprint(p)
        ilanlar.append(p)
    store.upsert(ilanlar)
    store.close()

    # sabit tarih yazmak testi takvime bagliyordu: first_seen_at hep "bugun"
    baslangic = (bugun - timedelta(days=7)).date().isoformat()
    bitis = bugun.date().isoformat()
    response = client.get(f"/?date_from={baslangic}&date_to={bitis}"
                          "&sort=budget&date_field=seen&min_score=0")
    assert response.status_code == 200
    for parca in (f"date_from={baslangic}", f"date_to={bitis}", "sort=budget", "date_field=seen"):
        assert parca in response.text, parca


def _ulke_ilanlari(tmp_path, kayitlar):
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    ilanlar = []
    for i, (ulke, konum) in enumerate(kayitlar):
        p = Project(source="test", source_id=str(i), url=f"http://x/{i}",
                    title=f"SAP ABAP {i}", company=f"F{i}", country=ulke, location=konum)
        p.fingerprint = fingerprint(p)
        ilanlar.append(p)
    store.upsert(ilanlar)
    store.close()


def test_ulke_filtresi_coklu_secim(client, tmp_path):
    """Tekrarli country parametresi birden fazla ulkeyi birlestirmeli."""
    _ulke_ilanlari(tmp_path, [("Germany", "Berlin"), ("Deutschland", "Köln"),
                              ("Österreich", "Wien"), ("Casablanca", "Casablanca")])

    tek = client.get("/?country=DE&min_score=0")
    assert tek.status_code == 200
    assert "SAP ABAP 0" in tek.text and "SAP ABAP 1" in tek.text
    assert "SAP ABAP 2" not in tek.text

    coklu = client.get("/?country=DE&country=AT&min_score=0")
    assert "SAP ABAP 2" in coklu.text
    assert "SAP ABAP 3" not in coklu.text          # cozulemeyen kayit gelmemeli


def test_ulke_kutusu_kanonik_adlari_gosterir(client, tmp_path):
    """Datalist ham 'TX'/'Remote' degerlerini listeliyordu; artik ulke adi + sayi."""
    _ulke_ilanlari(tmp_path, [("TX", "Austin, TX"), ("Remote", "Remote")])

    response = client.get("/?min_score=0")
    assert 'value="US"' in response.text
    assert "Amerika Birleşik Devletleri" in response.text
    assert "Ülke belirsiz" in response.text        # cozulemeyenler icin ayri kova
    assert 'list="country-list"' not in response.text


def test_ulke_secimi_sayfalamada_korunur(client, tmp_path):
    client.post("/ayarlar/tarama", data={"page_size": "5"}, follow_redirects=False)
    _ulke_ilanlari(tmp_path, [("Germany", "Berlin")] * 12)

    response = client.get("/?country=DE&min_score=0")
    assert response.status_code == 200
    assert "country=DE" in response.text


def test_eski_serbest_metin_ulke_filtresi_calisir(client, tmp_path):
    """Profillerde kayitli eski degerler (kod degil) sessizce sonucsuz kalmamali."""
    _ulke_ilanlari(tmp_path, [("Germany", "Berlin"), ("Österreich", "Wien")])

    response = client.get("/?country=Germany&min_score=0")
    assert "SAP ABAP 0" in response.text
    assert "SAP ABAP 1" not in response.text


def test_tarih_araligi_gun_filtresini_ezer(client):
    """Ikisi ayni ekseni filtreliyor; aralik seciliyse o kazanir."""
    response = client.get("/?days=7&date_from=2026-08-01")
    assert response.status_code == 200
    # "Yayin tarihi" secimi sifirlanmis olmali (7 secili kalmamali)
    assert '<option value="7" selected>' not in response.text
    assert 'value="2026-08-01"' in response.text


# --- zamanlayici ------------------------------------------------------
def test_reconfigure_thread_baslatmadan_uygular():
    scanner = BackgroundScanner(interval_minutes=30, quick_minutes=10)
    scanner.reconfigure(interval_minutes=60, quick_minutes=15, verify_limit=99)

    assert scanner.state.interval_minutes == 60
    assert scanner.state.quick_minutes == 15
    assert scanner.interval == 3600
    assert scanner.full_every == 4
    assert scanner.verify_limit == 99


def test_reconfigure_alt_sinirlari_zorlar():
    scanner = BackgroundScanner(interval_minutes=30, quick_minutes=10)
    scanner.reconfigure(interval_minutes=1, quick_minutes=1)
    assert scanner.state.interval_minutes == 5
    assert scanner.state.quick_minutes == 2


def test_duraklatma_thread_durdurmaz():
    """Masaustu uygulamasi zamanlayiciyi kendi yonetiyor; panel onu durdurmamali."""
    scanner = BackgroundScanner()
    scanner.set_enabled(False)
    assert scanner.state.enabled is False
    assert scanner._active.is_set() is False

    # "Şimdi tara" duraklatilmisken de bir tur istemeli
    assert scanner.trigger() is True
    assert scanner._force.is_set() is True

    scanner.set_enabled(True)
    assert scanner.state.enabled is True
    assert scanner._active.is_set() is True


# --- temizlik turu (Listeyi kontrol et) -----------------------------------

def test_temizlik_yetkisiz_kullaniciya_kapali(anon_client):
    assert anon_client.post("/api/temizlik", json={}, follow_redirects=False).status_code in (302, 303, 403)


def test_temizlik_baslar_ve_durumda_gorunur(client):
    response = client.post("/api/temizlik", json={"min_score": "0"})
    assert response.status_code == 200
    assert response.json()["ok"] is True

    durum = client.get("/api/durum").json()
    assert "sweep" in durum
    assert durum["sweep"]["phase"] in ("running", "done")


def test_temizlik_ikinci_istek_reddedilir(client, monkeypatch):
    """Ayni anda iki tur ag trafigini ikiye katlar; ikincisi 409 almali."""
    import scanner.web.app as app_mod

    # Tur bitmesin diye kontrol askida kalsin
    engel = __import__("threading").Event()

    def bekle(self, store, client_, rows, workers, profile_id, config):
        engel.wait(timeout=5)

    monkeypatch.setattr(app_mod.LinkSweeper, "_sweep", bekle)
    try:
        assert client.post("/api/temizlik", json={}).status_code == 200
        ikinci = client.post("/api/temizlik", json={})
        assert ikinci.status_code == 409
        assert "sürüyor" in ikinci.json()["message"]
    finally:
        engel.set()

def test_temizlik_ekrandaki_filtreyi_kullanir(client, monkeypatch):
    """Tur, adres cubugundaki filtrenin AYNISINI tarar.

    Onay penceresi "ekranda gordugun N ilan kontrol edilecek" diyor; sunucu
    baska bir kume tararsa bu yalan olur. Kullanici tam bunu yasadi: pencere
    1805 derken listesinde 346 ilan vardi.
    """
    import scanner.web.app as app_mod

    yakalanan: dict = {}

    def yakala(self, filters, by="", config=None, full=False):
        yakalanan.update(filters)
        yakalanan["_full"] = full
        return True

    monkeypatch.setattr(app_mod.LinkSweeper, "start", yakala)
    client.post("/api/temizlik",
                json={"query": "?min_score=40&source=jooble&country=DE&country=AT"})

    assert yakalanan["min_score"] == 40
    assert yakalanan["source"] == "jooble"
    assert yakalanan["country"] == ["DE", "AT"]      # tekrarli parametre korunur
    assert yakalanan["_full"] is True                 # kume tam taranir, kirpilmaz


def test_temizlik_bos_sorgu_ile_profil_varsayilanina_duser(client, monkeypatch):
    """Gövdesiz/bos istek (eski istemci) 422 vermemeli, tur yine calismali."""
    import scanner.web.app as app_mod

    yakalanan: dict = {}
    monkeypatch.setattr(app_mod.LinkSweeper, "start",
                        lambda self, filters, by="", config=None, full=False:
                        (yakalanan.update(filters), True)[1])
    cevap = client.post("/api/temizlik")
    assert cevap.status_code == 200
    assert yakalanan["closed_only"] is False          # yalnizca aktif ilanlar
    assert yakalanan["source"] is None                # kaynak filtresi yok
    assert yakalanan["search"] is None
    assert yakalanan["source"] is None
    # `closed` verilmez: Storage varsayilani yalnizca aktif ilanlari getirir.
    assert not yakalanan.get("closed_only")


def test_temizlik_en_eski_kontrol_edilenden_baslar():
    """`stale` sirasi ORDERS icinde tanimli olmali.

    Tur basina 500 ilan sinirli: varsayilan puan siralamasiyla her tur ayni
    en yuksek puanli ilanlar taranir, listenin kuyruguna hic sira gelmezdi.
    """
    from scanner.storage import Storage

    assert "stale" in Storage.ORDERS
    sira = Storage.ORDERS["stale"]
    assert sira.startswith("missing_streak DESC")   # olme ihtimali yuksek olan once
    assert "verified_at" in sira                    # sonra en uzun suredir bakilmayan

# --- ilan ozeti (merkez modal, /api/ozet) ----------------------------------

def _tek_ilan(tmp_path, description="", title="SAP ABAP Test", score=0):
    """Tek ilanli bir DB kurar, fingerprint'i dondurur - /api/ozet testleri icin."""
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    p = Project(source="test", source_id="1", url="http://x/1", title=title,
               description=description, keywords_hit=["abap", "sap", "remote"],
               score=score)
    p.fingerprint = fingerprint(p)
    store.upsert([p])
    store.close()
    return p.fingerprint


def test_ozet_bilinmeyen_fingerprint_404(client):
    assert client.get("/api/ozet?fingerprint=yok").status_code == 404


def test_ozet_yaniti_zorunlu_alanlari_icerir(client, tmp_path, monkeypatch):
    """`requirements` yalnizca must_any ile kesisen kelimeleri icermeli."""
    import scanner.web.app as app_mod

    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP remote consultant.")
    monkeypatch.setattr(app_mod, "detect_language", lambda text: "en")

    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert veri["ok"] is True
    assert veri["title"] == "SAP ABAP Test"
    assert veri["summary"].startswith("We need an SAP ABAP")
    assert set(veri["requirements"]) == {"abap", "sap"}   # "remote" must_any'de degil
    assert veri["translated"] is False


def test_ozet_ingilizce_ilan_cevrilmez(client, tmp_path, monkeypatch):
    """Ceviri servisi hic cagrilmamali: kota bosa harcanmasin."""
    import scanner.translate as translate_mod

    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP remote consultant.")
    cagrildi = []
    monkeypatch.setattr(translate_mod, "translate",
                        lambda *a, **k: cagrildi.append(1) or "should not happen")

    client.get(f"/api/ozet?fingerprint={fp}")
    assert cagrildi == []


def test_ozet_ceviri_cache_lenir(client, tmp_path, monkeypatch):
    """Ikinci istekte ceviri servisine tekrar gidilmemeli - DB'den okunmali."""
    import scanner.web.app as app_mod

    fp = _tek_ilan(tmp_path, description="Wir suchen einen SAP ABAP Entwickler.")
    monkeypatch.setattr(app_mod, "detect_language", lambda text: "de")
    cagri_sayisi = [0]

    def sahte_ceviri(text, source="", target="en", email=""):
        cagri_sayisi[0] += 1
        return "We are looking for an SAP ABAP developer."

    monkeypatch.setattr(app_mod, "translate_text", sahte_ceviri)

    ilk = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert ilk["translated"] is True
    assert ilk["lang"] == "de"
    assert ilk["summary"] == "We are looking for an SAP ABAP developer."

    ikinci = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert ikinci["summary"] == ilk["summary"]
    assert cagri_sayisi[0] == 1        # servise ikinci kez gidilmedi


def test_ozet_bos_aciklama(client, tmp_path):
    fp = _tek_ilan(tmp_path, description="")
    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert veri["ok"] is True
    assert veri["summary"] == ""
    assert veri["translated"] is False

# --- yapay zeka ozeti -------------------------------------------------------

def _ai_kur(monkeypatch, client, sonuc, *, esik=30, acik=True):
    """AI'yi yapilandirir ve sahte ozetleyiciyi baglar; cagri sayacini dondurur."""
    import scanner.web.app as app_mod

    monkeypatch.setenv("GEMINI_API_KEY", "test-anahtari")
    sayac = {"n": 0}

    def sahte_ozetle(baslik, aciklama, *, anahtar, model, zaman_asimi):
        sayac["n"] += 1
        return sonuc

    monkeypatch.setattr(app_mod.ai_mod, "ozetle", sahte_ozetle)
    client.post("/ayarlar/yapayzeka", data={
        "ai_enabled": "1" if acik else "0", "ai_min_score": str(esik),
        "ai_model": "test-model", "ai_timeout": "30",
    }, follow_redirects=False)
    return sayac


AI_SONUC = {"summary": "Remote SAP ABAP contract, 6 months.",
            "must_haves": ["5+ years ABAP", "Fluent English"]}


def test_ai_esik_ustundeki_ilani_ozetler(client, tmp_path, monkeypatch):
    fp = _tek_ilan(tmp_path, description="Wir suchen SAP ABAP Entwickler.", score=45)
    sayac = _ai_kur(monkeypatch, client, AI_SONUC, esik=30)

    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert veri["kaynak"] == "ai"
    assert veri["summary"] == AI_SONUC["summary"]
    assert veri["requirements"] == AI_SONUC["must_haves"]
    assert sayac["n"] == 1


def test_ai_esik_altindaki_ilani_gondermez(client, tmp_path, monkeypatch):
    """Token disiplini: dusuk skorlu ilan yapay zekaya HIC gitmemeli."""
    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP consultant.", score=10)
    sayac = _ai_kur(monkeypatch, client, AI_SONUC, esik=30)

    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert sayac["n"] == 0
    assert veri["kaynak"] == "kelime"
    assert set(veri["requirements"]) == {"abap", "sap"}


def test_ai_kapaliyken_gonderilmez(client, tmp_path, monkeypatch):
    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP consultant.", score=90)
    sayac = _ai_kur(monkeypatch, client, AI_SONUC, esik=0, acik=False)

    assert client.get(f"/api/ozet?fingerprint={fp}").json()["kaynak"] == "kelime"
    assert sayac["n"] == 0


def test_ai_ozeti_onbelleklenir(client, tmp_path, monkeypatch):
    """Ikinci istekte tekrar token harcanmamali - DB'den okunmali."""
    fp = _tek_ilan(tmp_path, description="Wir suchen SAP ABAP Entwickler.", score=45)
    sayac = _ai_kur(monkeypatch, client, AI_SONUC, esik=30)

    ilk = client.get(f"/api/ozet?fingerprint={fp}").json()
    ikinci = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert ikinci["summary"] == ilk["summary"]
    assert ikinci["requirements"] == ilk["requirements"]
    assert sayac["n"] == 1                  # servise ikinci kez gidilmedi


def test_ai_hata_verince_eski_davranisa_duser(client, tmp_path, monkeypatch):
    """Kota dolmus/anahtar bozuksa panel bos kalmamali."""
    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP consultant.", score=90)
    sayac = _ai_kur(monkeypatch, client, None, esik=0)      # ozetle() None doner

    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert sayac["n"] == 1
    assert veri["ok"] is True
    assert veri["kaynak"] == "kelime"
    assert veri["summary"].startswith("We need an SAP ABAP")


def test_ai_maliyet_kaydi_ve_ekrani(client, tmp_path, monkeypatch):
    """Basarili ozet maliyet gecmisine yazilir; /ayarlar/maliyet modeli ve tutari gosterir."""
    fp = _tek_ilan(tmp_path, description="Wir suchen SAP ABAP Entwickler.", score=45)
    sonuc = {**AI_SONUC, "prompt_tokens": 1300, "output_tokens": 200}
    _ai_kur(monkeypatch, client, sonuc, esik=30)

    client.get(f"/api/ozet?fingerprint={fp}")
    sayfa = client.get("/ayarlar/maliyet")
    assert sayfa.status_code == 200
    assert "test-model" in sayfa.text
    assert "1.500 tok" in sayfa.text          # 1300 + 200, nokta ayrac


def test_ai_anahtarsizken_gonderilmez(client, tmp_path, monkeypatch):
    import scanner.web.app as app_mod

    fp = _tek_ilan(tmp_path, description="We need an SAP ABAP consultant.", score=90)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    sayac = {"n": 0}
    monkeypatch.setattr(app_mod.ai_mod, "ozetle",
                        lambda *a, **k: sayac.__setitem__("n", sayac["n"] + 1))

    assert client.get(f"/api/ozet?fingerprint={fp}").json()["kaynak"] == "kelime"
    assert sayac["n"] == 0


def test_ai_ayarlari_overlaye_yazilir(client, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "mevcut")
    client.post("/ayarlar/yapayzeka", data={
        "ai_enabled": "1", "ai_min_score": "-5",      # negatif deger kirpilmali
        "ai_model": "  ", "ai_timeout": "1",           # bos model -> varsayilan
    }, follow_redirects=False)

    data = config_mod.load_overlay(config_mod.overlay_path())
    assert data["ai"]["min_score"] == 0                # max(0, -5)
    assert data["ai"]["timeout"] == 5                  # max(5, 1)
    # Model overlay'e YAZILMAZ: bos birakilinca varsayilana dusuyor, varsayilan da
    # config.yaml'daki degerle ayni - `prune_unchanged` onu eliyor. Boylece ileride
    # varsayilan model degisirse kullaniciya kendiliginden ulasir.
    assert "model" not in data["ai"]

    # Farkli bir model verilirse overlay'e yazilmali:
    client.post("/ayarlar/yapayzeka", data={
        "ai_enabled": "1", "ai_min_score": "30",
        "ai_model": "gemini-2.5-flash-lite", "ai_timeout": "30",
    }, follow_redirects=False)
    data = config_mod.load_overlay(config_mod.overlay_path())
    assert data["ai"]["model"] == "gemini-2.5-flash-lite"


def test_ai_anahtari_bos_birakilirsa_korunur(client, monkeypatch, tmp_path):
    """Kaydet'e basarken alan bos ise mevcut anahtar silinmemeli."""
    import scanner.config as config_mod2

    monkeypatch.setenv("GEMINI_API_KEY", "eski-anahtar")
    yazilanlar = []
    monkeypatch.setattr(config_mod2, "save_env", lambda v: yazilanlar.append(v))
    import scanner.web.app as app_mod
    monkeypatch.setattr(app_mod, "save_env", lambda v: yazilanlar.append(v))

    client.post("/ayarlar/yapayzeka", data={
        "ai_enabled": "1", "ai_min_score": "30", "ai_model": "m", "ai_timeout": "30",
    }, follow_redirects=False)
    assert yazilanlar == []                 # .env'e hic dokunulmadi


# --- kapanan bildirimleri temizle -----------------------------------------

def test_bildirim_sil_sadece_kapananlar(client, tmp_path):
    from scanner.storage import Storage
    store = Storage(tmp_path / "data" / "projects.db")
    store.add_notification("new", "Yeni ilan")
    store.add_notification("closed", "Kapanan ilan A")
    store.add_notification("closed", "Kapanan ilan B")
    store.close()

    client.post("/bildirimler/sil", data={"hepsi": "1", "kind": "closed"},
                follow_redirects=False)

    store = Storage(tmp_path / "data" / "projects.db")
    kalan = store.notifications()
    store.close()
    kinds = sorted(n["kind"] for n in kalan)
    assert kinds == ["new"]


# --- basvuru takip ekrani ------------------------------------------------

def test_basvurular_ekrani_ve_durum(client, tmp_path):
    fp = _tek_ilan(tmp_path, title="SAP ABAP Başvuru")
    client.post("/status", data={"fingerprint": fp, "status": "applied", "back": "/"},
                follow_redirects=False)

    sayfa = client.get("/basvurular")
    assert sayfa.status_code == 200
    assert "SAP ABAP Başvuru" in sayfa.text

    # süreç durumunu güncelle
    client.post("/status", data={"fingerprint": fp, "status": "won_active", "back": "/basvurular"},
                follow_redirects=False)
    from scanner.storage import Storage
    store = Storage(tmp_path / "data" / "projects.db")
    assert store.status_of(fp) == "won_active"
    assert [r["mark"] for r in store.applications()] == ["won_active"]
    store.close()


def test_ilan_word_ciktisi_gecerli_docx(client, tmp_path):
    import io
    import zipfile
    from xml.dom.minidom import parseString

    fp = _tek_ilan(tmp_path, title="SAP ABAP Word Testi", score=77)
    r = client.get(f"/ilan/{fp}/word")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "SAP-ABAP-Word-Testi.docx" in r.headers["content-disposition"]

    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert "word/document.xml" in z.namelist()
    doc = z.read("word/document.xml").decode("utf-8")
    parseString(doc)                       # iyi bicimli XML
    assert "SAP ABAP Word Testi" in doc
    assert "Künye" in doc and "77" in doc


def test_ilan_word_bilinmeyen_fingerprint_404(client):
    assert client.get("/ilan/yok/word").status_code == 404


# --- "Dışa aktar" ekrani (coklu secim -> xlsx) --------------------------

def test_disari_aktar_ekrani_filtreli_liste(client, tmp_path):
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    for i, (t, src) in enumerate([("SAP ABAP Remote", "jooble"), ("Java Dev", "reed")]):
        p = Project(source=src, source_id=str(i), url=f"http://x/{i}", title=t, score=70)
        p.fingerprint = fingerprint(p)
        store.upsert([p])
    store.close()

    hepsi = client.get("/disari-aktar?min_score=0")
    assert hepsi.status_code == 200
    assert "SAP ABAP Remote" in hepsi.text and "Java Dev" in hepsi.text
    # "tümünü seç" kutusu YOK
    assert 'id="select-all"' not in hepsi.text

    suzulmus = client.get("/disari-aktar?min_score=0&q=abap").text
    assert "SAP ABAP Remote" in suzulmus and "Java Dev" not in suzulmus


def test_disari_aktar_secilenleri_xlsx_verir(client, tmp_path):
    import io
    import zipfile
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    fps = []
    for i in range(3):
        p = Project(source="jooble", source_id=str(i), url=f"http://x/{i}",
                    title=f"İlan {i}", score=60, keywords_hit=["abap"])
        p.fingerprint = fingerprint(p)
        fps.append(p.fingerprint)
        store.upsert([p])
    store.close()

    r = client.post("/disari-aktar", data={"fp": [fps[0], fps[2]]})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "secili-ilanlar" in r.headers["content-disposition"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    basliklar = [c.value for c in ws[1]]
    assert "Güncellik" in basliklar and "Skor" in basliklar and "Önemli hususlar" in basliklar
    assert ws.max_row == 3                              # 2 seçilen + başlık satırı


def test_disari_aktar_bos_secim_geri_doner(client):
    r = client.post("/disari-aktar", data={}, follow_redirects=False)
    assert r.status_code == 303


def test_basvurulan_ilan_ana_listeden_kalkar(client, tmp_path):
    """Basvuru surecindeki ilan varsayilan listede digerleriyle karismaz."""
    fp = _tek_ilan(tmp_path, title="SAP ABAP Gizlenecek", score=80)
    assert fp in client.get("/?min_score=0").text

    client.post("/status", data={"fingerprint": fp, "status": "applied"}, follow_redirects=False)
    assert "SAP ABAP Gizlenecek" not in client.get("/?min_score=0").text

    # ama Durum filtresiyle acikca istenirse yine gorunur
    assert "SAP ABAP Gizlenecek" in client.get("/?min_score=0&status=applied").text


def test_basvurular_durum_grubu_filtreler(client, tmp_path):
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    fps = []
    for i, baslik in enumerate(["A", "B", "C"]):
        p = Project(source="test", source_id=str(i), url=f"http://x/{i}", title=f"İlan {baslik}")
        p.fingerprint = fingerprint(p)
        fps.append(p.fingerprint)
        store.upsert([p])
    store.close()

    client.post("/status", data={"fingerprint": fps[0], "status": "applied"}, follow_redirects=False)
    client.post("/status", data={"fingerprint": fps[1], "status": "won_active"}, follow_redirects=False)
    client.post("/status", data={"fingerprint": fps[2], "status": "won_done"}, follow_redirects=False)

    tumu = client.get("/basvurular").text
    assert "İlan A" in tumu and "İlan B" in tumu and "İlan C" in tumu

    olumlu = client.get("/basvurular?durum=won").text
    assert "İlan A" not in olumlu
    assert "İlan B" in olumlu and "İlan C" in olumlu

    basvuruldu = client.get("/basvurular?durum=applied").text
    assert "İlan A" in basvuruldu
    assert "İlan B" not in basvuruldu and "İlan C" not in basvuruldu

    # tek tek olumlu alt durumlari da ayri filtrelenebilmeli
    sadece_aktif = client.get("/basvurular?durum=won_active").text
    assert "İlan B" in sadece_aktif
    assert "İlan A" not in sadece_aktif and "İlan C" not in sadece_aktif


def test_basvurular_arama_kutusu_filtreler(client, tmp_path):
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage

    store = Storage(tmp_path / "data" / "projects.db")
    fps = []
    for baslik, firma in [("SAP ABAP Danışmanı", "Acme"), ("Java Geliştirici", "Beta")]:
        p = Project(source="test", source_id=baslik, url=f"http://x/{baslik}",
                    title=baslik, company=firma)
        p.fingerprint = fingerprint(p)
        fps.append(p.fingerprint)
        store.upsert([p])
    store.close()
    for fp in fps:
        client.post("/status", data={"fingerprint": fp, "status": "applied"}, follow_redirects=False)

    veri = client.get("/basvurular?q=abap").text
    assert "SAP ABAP Danışmanı" in veri
    assert "Java Geliştirici" not in veri


def test_basvuru_kaldir_ana_listeye_geri_dondurur(client, tmp_path):
    """Yanlislikla isaretlenen basvuru 'Kaldir' ile normal ilana doner."""
    fp = _tek_ilan(tmp_path, title="Yanlış Tıklanan İlan", score=80)
    client.post("/status", data={"fingerprint": fp, "status": "applied"}, follow_redirects=False)
    assert "Yanlış Tıklanan İlan" not in client.get("/?min_score=0").text
    assert "Yanlış Tıklanan İlan" in client.get("/basvurular").text

    client.post("/status", data={"fingerprint": fp, "status": "new"}, follow_redirects=False)

    assert "Yanlış Tıklanan İlan" not in client.get("/basvurular").text
    assert "Yanlış Tıklanan İlan" in client.get("/?min_score=0").text


def test_basvuru_sil_ilani_tamamen_kaldirir(client, tmp_path):
    fp = _tek_ilan(tmp_path, title="Silinecek İlan")
    client.post("/status", data={"fingerprint": fp, "status": "applied"}, follow_redirects=False)
    client.post("/basvuru/sil", data={"fingerprint": fp}, follow_redirects=False)

    from scanner.storage import Storage
    store = Storage(tmp_path / "data" / "projects.db")
    assert store.get_by_fingerprint(fp) is None
    store.close()


# --- freelancermap özet zenginleştirme ----------------------------------

def test_freelancermap_ozet_detay_sayfasindan_zenginlesir(client, tmp_path, monkeypatch):
    from scanner.dedupe import fingerprint
    from scanner.models import Project
    from scanner.storage import Storage
    import scanner.web.app as app_mod

    store = Storage(tmp_path / "data" / "projects.db")
    p = Project(source="freelancermap", source_id="9", url="http://fm/9",
                title="SAP ABAP Contract", description="SAP | ABAP", score=90)
    p.fingerprint = fingerprint(p)
    store.upsert([p])
    store.close()
    fp = p.fingerprint

    uzun = ("We are looking for a senior SAP ABAP consultant for a fully remote "
            "six month contract. Strong S/4HANA migration background required.")
    from scanner.sources import freelancermap as fm
    monkeypatch.setattr(fm, "detail_description", lambda url, **k: uzun)
    # AI kapalı: eski davranış → ham açıklama özet olarak döner
    monkeypatch.setattr(app_mod.ai_mod, "ozetle", lambda *a, **k: None)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    veri = client.get(f"/api/ozet?fingerprint={fp}").json()
    assert veri["ok"] is True
    assert "S/4HANA migration" in veri["summary"]

    store = Storage(tmp_path / "data" / "projects.db")
    assert "S/4HANA migration" in store.get_by_fingerprint(fp)["description"]
    store.close()
